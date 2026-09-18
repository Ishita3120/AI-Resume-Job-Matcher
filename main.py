import os
import json
from pathlib import Path
from typing import Optional

from dotenv import load_dotenv
from pypdf import PdfReader
from pydantic import BaseModel, Field
from groq import Groq
from docx import Document


# ============================================================
# 1. RESUME FOLDER
# ============================================================

resume_folder = Path("resumes")


# ============================================================
# 2. LOAD ENVIRONMENT VARIABLES
# ============================================================

load_dotenv()

api_key = os.getenv("GROK_API_KEY")

if not api_key:
    raise ValueError(
        "GROK API key not found. Please add it to your .env file."
    )


# ============================================================
# 3. GROQ CLIENT
# ============================================================

client = Groq(api_key=api_key)

model = "openai/gpt-oss-120b"


# ============================================================
# 4. JOB DESCRIPTION
# ============================================================

job_description = """
Software Engineer Intern

Requirements:
- Strong knowledge of Python
- Data Structures and Algorithms
- SQL
- React
- REST APIs
- Git
- Good problem solving skills
- Knowledge of Machine Learning is a plus
"""


# ============================================================
# 5. RESUME PYDANTIC MODEL
# ============================================================

class Resume(BaseModel):
    name: str
    email: str
    phone: str
    skills: list[str] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    experience: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    achievements: Optional[list[str]] = None
    hobbies: Optional[list[str]] = None


# ============================================================
# 6. JOB MATCH PYDANTIC MODEL
# ============================================================

class JobMatch(BaseModel):
    match_percentage: int
    matched_skills: list[str] = Field(default_factory=list)
    missing_skills: list[str] = Field(default_factory=list)
    explanation: str


# ============================================================
# 7. CREATE JSON SCHEMAS
# ============================================================

resume_schema = Resume.model_json_schema()

match_schema = JobMatch.model_json_schema()


# ============================================================
# 8. CHECK RESUME FOLDER
# ============================================================

if not resume_folder.exists():
    raise FileNotFoundError(
        "Resumes folder not found. "
        "Please create a 'resumes' folder inside the project."
    )


# ============================================================
# 9. FIND ALL PDF AND DOCX RESUMES
# ============================================================

resume_files = (
    list(resume_folder.glob("*.pdf"))
    + list(resume_folder.glob("*.docx"))
)

if not resume_files:
    raise FileNotFoundError(
        "No PDF or DOCX resumes found inside the resumes folder."
    )

print(f"Found {len(resume_files)} resume(s).")


# ============================================================
# 10. TEXT EXTRACTION FUNCTION
# ============================================================

def extract_text(file_path: Path) -> str:

    extension = file_path.suffix.lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        reader = PdfReader(file_path)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text


    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        document = Document(file_path)

        text = ""

        # Extract paragraphs
        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text += paragraph.text + "\n"


        # Extract tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    row_text.append(cell.text.strip())

                text += " | ".join(row_text) + "\n"


        return text


    # --------------------------------------------------------
    # Unsupported format
    # --------------------------------------------------------

    else:

        raise ValueError(
            "Unsupported file format. "
            "Only PDF and DOCX are supported."
        )


# ============================================================
# 11. PROCESS EACH RESUME
# ============================================================

for resume_file in resume_files:

    print("\n")
    print("=" * 70)
    print(f"PROCESSING: {resume_file.name}")
    print("=" * 70)


    # --------------------------------------------------------
    # Extract resume text
    # --------------------------------------------------------

    resume_text = extract_text(resume_file)

    if not resume_text.strip():

        print("Could not extract any text from this resume.")

        continue


    # ========================================================
    # PART A — EXTRACT RESUME INFORMATION
    # ========================================================

    resume_response_format = {

        "type": "json_schema",

        "json_schema": {

            "name": "resume",

            "schema": resume_schema
        }
    }


    resume_messages = [

        {
            "role": "system",

            "content": (
                "Extract information from the resume according "
                "to the provided schema. "
                "Do not invent information that is not present "
                "in the resume."
            )
        },

        {
            "role": "user",

            "content": resume_text
        }
    ]


    resume_response = client.chat.completions.create(

        model=model,

        messages=resume_messages,

        response_format=resume_response_format
    )


    resume_answer = resume_response.choices[0].message.content

    resume_data = json.loads(resume_answer)

    resume = Resume.model_validate(resume_data)


    print("\n--- EXTRACTED RESUME ---")

    print(resume)


    # ========================================================
    # PART B — MATCH RESUME WITH JOB DESCRIPTION
    # ========================================================

    match_response_format = {

        "type": "json_schema",

        "json_schema": {

            "name": "job_match",

            "schema": match_schema
        }
    }


    match_messages = [

        {
            "role": "system",

            "content": """
You are a resume and job-description matching system.

Compare the candidate's resume with the given job description.

Give a match percentage from 0 to 100 based on how well
the candidate's documented skills, education, experience,
and projects align with the job requirements.

Do not invent skills, education, experience, or projects.

Identify:
1. Skills that match the job description.
2. Important skills or requirements missing from the resume.
3. An overall match percentage.
4. A short explanation for the percentage.

The match percentage represents requirement alignment only.
It is not a prediction of hiring or selection.
"""
        },

        {
            "role": "user",

            "content": f"""
JOB DESCRIPTION:

{job_description}


CANDIDATE RESUME:

{resume_text}
"""
        }
    ]


    match_response = client.chat.completions.create(

        model=model,

        messages=match_messages,

        response_format=match_response_format
    )


    match_answer = match_response.choices[0].message.content

    match_data = json.loads(match_answer)

    match = JobMatch.model_validate(match_data)


    # ========================================================
    # PART C — DISPLAY MATCH RESULT
    # ========================================================

    print("\n--- JOB MATCH RESULT ---")

    print(f"Match Percentage: {match.match_percentage}%")

    print("\nMatched Skills:")

    for skill in match.matched_skills:

        print(f"  ✓ {skill}")


    print("\nMissing Skills:")

    for skill in match.missing_skills:

        print(f"  ✗ {skill}")


    print("\nExplanation:")

    print(match.explanation)